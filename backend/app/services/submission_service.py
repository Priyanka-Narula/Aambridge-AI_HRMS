import io
import json
import uuid
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from typing import Any

from fastapi import HTTPException, status
from sqlalchemy.orm import Session, joinedload

from app.core.security import ROLE_OWNER, ROLE_RECRUITER, normalize_role
from app.models.candidate import Candidate
from app.models.job_requirement import JobRequirement
from app.models.pipeline import CandidateApplication, PipelineStage
from app.models.user_access import Client, Recruiter, User


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _load_job_requirement(db: Session, job_req_id: uuid.UUID) -> JobRequirement:
    jr = (
        db.query(JobRequirement)
        .options(joinedload(JobRequirement.client))
        .filter(JobRequirement.id == job_req_id)
        .first()
    )
    if not jr:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Job requirement not found")
    return jr


def _load_candidate(db: Session, candidate_id: uuid.UUID) -> Candidate:
    candidate = db.query(Candidate).filter(Candidate.id == candidate_id).first()
    if not candidate:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Candidate not found")
    return candidate


def _load_application(db: Session, app_id: uuid.UUID) -> CandidateApplication:
    app = (
        db.query(CandidateApplication)
        .options(
            joinedload(CandidateApplication.current_stage_rel),
            joinedload(CandidateApplication.submitter),
        )
        .filter(CandidateApplication.id == app_id)
        .first()
    )
    if not app:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Submission not found")
    return app


def _load_application_with_full_relations(db: Session, app_id: uuid.UUID) -> CandidateApplication:
    return (
        db.query(CandidateApplication)
        .options(
            joinedload(CandidateApplication.current_stage_rel),
            joinedload(CandidateApplication.submitter),
        )
        .filter(CandidateApplication.id == app_id)
        .one()
    )


def _get_first_stage(db: Session) -> PipelineStage | None:
    return db.query(PipelineStage).order_by(PipelineStage.order_no).first()


def _get_recruiter(db: Session, user: User) -> Recruiter | None:
    if user.recruiter:
        return user.recruiter
    return db.query(Recruiter).filter(Recruiter.user_id == user.id).first()


_DEFAULT_SUBMISSION_FIELDS: list[dict[str, str]] = [
    {"field": "First Name"},
    {"field": "Last Name"},
    {"field": "Email"},
    {"field": "Phone"},
    {"field": "Current Company"},
    {"field": "Current Designation"},
    {"field": "Total Experience (Years)"},
    {"field": "Notice Period"},
    {"field": "Current CTC"},
    {"field": "Expected CTC"},
    {"field": "Visa Status"},
    {"field": "Current Location"},
]

# Map common template labels → Candidate model attributes (case-insensitive).
_CANDIDATE_LABEL_ATTRS: dict[str, str] = {
    "first name": "first_name",
    "last name": "last_name",
    "email": "email",
    "phone": "phone",
    "mobile": "phone",
    "current company": "current_company",
    "company": "current_company",
    "current designation": "current_designation",
    "designation": "current_designation",
    "total experience (years)": "total_experience_years",
    "total experience": "total_experience_years",
    "experience": "total_experience_years",
    "notice period": "notice_period",
    "current ctc": "current_ctc",
    "expected ctc": "expected_ctc",
    "visa status": "visa_status",
    "current location": "current_location",
    "location": "current_location",
    "nationality": "nationality",
    "linkedin": "linkedin_url",
    "linkedin url": "linkedin_url",
}


def _parse_submission_fields(raw: str | None) -> list[dict]:
    if not raw:
        return list(_DEFAULT_SUBMISSION_FIELDS)
    try:
        fields = json.loads(raw)
    except (json.JSONDecodeError, TypeError):
        return list(_DEFAULT_SUBMISSION_FIELDS)
    return fields if fields else list(_DEFAULT_SUBMISSION_FIELDS)


def _parse_stored_submission_data(raw: str | None) -> dict[str, Any]:
    if not raw:
        return {}
    try:
        data = json.loads(raw)
    except (json.JSONDecodeError, TypeError):
        return {}
    return data if isinstance(data, dict) else {}


def _candidate_attr_value(candidate: Candidate | None, label: str) -> str | None:
    if not candidate or not label:
        return None
    attr = _CANDIDATE_LABEL_ATTRS.get(label.strip().lower())
    if not attr:
        return None
    value = getattr(candidate, attr, None)
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _cell_value(stored_data: dict[str, Any], label: str, candidate: Candidate | None) -> str | None:
    """Prefer recruiter-filled submission_data; fall back to Candidate fields."""
    if label in stored_data:
        raw = stored_data.get(label)
        if raw is not None and str(raw).strip():
            return str(raw).strip()
    # Case-insensitive key match in stored data
    label_l = label.strip().lower()
    for key, raw in stored_data.items():
        if str(key).strip().lower() == label_l and raw is not None and str(raw).strip():
            return str(raw).strip()
    return _candidate_attr_value(candidate, label)


# ---------------------------------------------------------------------------
# A. Mandatory fields check against submission_data
# ---------------------------------------------------------------------------

def check_mandatory_fields(
    submission_format_json: str | None,
    submission_data: dict[str, Any] | None,
) -> list[str]:
    """
    Check that every required field in the client's submission_format
    has a non-empty value in submission_data.

    submission_format_json: the JSON string stored on Client.submission_format
    submission_data: dict keyed by field label (e.g. {"First Name": "John", ...})

    Returns a list of missing required field labels.
    """
    if not submission_format_json:
        return []
    try:
        fields = json.loads(submission_format_json)
    except (json.JSONDecodeError, TypeError):
        return []

    data = submission_data or {}
    missing = []
    for field_def in fields:
        if not field_def.get("required", False):
            continue
        label = field_def.get("field", "")
        if not label:
            continue
        value = data.get(label)
        if value is None or (isinstance(value, str) and not value.strip()):
            missing.append(label)
    return missing


# ---------------------------------------------------------------------------
# B. Deduplication checks
# ---------------------------------------------------------------------------

def check_submission_eligibility(
    db: Session,
    candidate_id: uuid.UUID,
    job_requirement: JobRequirement,
) -> None:
    """Raise 409 if either dedup rule is violated."""
    now = datetime.now(timezone.utc)
    client_id = job_requirement.client_id

    # Rule A — same client, within 6 months
    cutoff_6mo = now - timedelta(days=182)
    existing_same_client = (
        db.query(CandidateApplication)
        .join(JobRequirement, CandidateApplication.job_requirement_id == JobRequirement.id)
        .filter(
            CandidateApplication.candidate_id == candidate_id,
            JobRequirement.client_id == client_id,
            CandidateApplication.submitted_at >= cutoff_6mo,
        )
        .first()
    )
    if existing_same_client:
        shared_date = (
            existing_same_client.submitted_at.strftime("%d %b %Y")
            if existing_same_client.submitted_at
            else "a recent date"
        )
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                f"This CV was already shared with {job_requirement.client.company_name} "
                f"on {shared_date}. Cannot re-share within 6 months."
            ),
        )

    # Rule B — different client, within 3 months, still in process
    cutoff_3mo = now - timedelta(days=91)
    existing_other = (
        db.query(CandidateApplication, Client)
        .join(JobRequirement, CandidateApplication.job_requirement_id == JobRequirement.id)
        .join(Client, JobRequirement.client_id == Client.id)
        .filter(
            CandidateApplication.candidate_id == candidate_id,
            JobRequirement.client_id != client_id,
            CandidateApplication.submitted_at >= cutoff_3mo,
            CandidateApplication.status.notin_(["rejected", "withdrawn"]),
        )
        .first()
    )
    if existing_other:
        _, other_client = existing_other
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                f"Candidate is already in process with {other_client.company_name}. "
                f"Cannot submit to another client within 3 months."
            ),
        )


# ---------------------------------------------------------------------------
# C. Submit candidate
# ---------------------------------------------------------------------------

def submit_candidate(
    db: Session,
    job_req_id: uuid.UUID,
    candidate_id: uuid.UUID,
    current_user: User,
    submission_data: dict[str, Any] | None = None,
) -> CandidateApplication:
    jr = _load_job_requirement(db, job_req_id)
    role = normalize_role(current_user.role.name)

    if role == ROLE_RECRUITER and jr.assigned_to != current_user.id:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="You are not assigned to this job requirement",
        )

    _load_candidate(db, candidate_id)

    # Mandatory fields check — against the recruiter-filled submission_data
    client_fmt = jr.client.submission_format if jr.client else None
    missing = check_mandatory_fields(client_fmt, submission_data)
    if missing:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Missing required fields for this client: {', '.join(missing)}",
        )

    # Dedup checks
    check_submission_eligibility(db, candidate_id, jr)

    # Prevent re-submission to same job
    already_submitted = (
        db.query(CandidateApplication)
        .filter(
            CandidateApplication.candidate_id == candidate_id,
            CandidateApplication.job_requirement_id == job_req_id,
        )
        .first()
    )
    if already_submitted:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Candidate is already submitted to this job requirement",
        )

    first_stage = _get_first_stage(db)
    recruiter = _get_recruiter(db, current_user) if role == ROLE_RECRUITER else None

    application = CandidateApplication(
        candidate_id=candidate_id,
        job_requirement_id=job_req_id,
        assigned_recruiter_id=recruiter.id if recruiter else None,
        current_stage=first_stage.id if first_stage else None,
        applied_date=datetime.now(timezone.utc).date(),
        status="active",
        submitted_by=current_user.id,
        submitted_at=datetime.now(timezone.utc),
        owner_status="pending_review",
        submission_data=json.dumps(submission_data) if submission_data else None,
    )
    db.add(application)
    db.commit()

    return _load_application_with_full_relations(db, application.id)


# ---------------------------------------------------------------------------
# D. List submissions for a job requirement
# ---------------------------------------------------------------------------

def list_submissions(
    db: Session,
    job_req_id: uuid.UUID,
    current_user: User,
) -> list[CandidateApplication]:
    jr = _load_job_requirement(db, job_req_id)
    role = normalize_role(current_user.role.name)

    if role == ROLE_RECRUITER and jr.assigned_to != current_user.id:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Insufficient permissions")

    return (
        db.query(CandidateApplication)
        .filter(CandidateApplication.job_requirement_id == job_req_id)
        .options(joinedload(CandidateApplication.current_stage_rel))
        .order_by(CandidateApplication.submitted_at.desc())
        .all()
    )


# ---------------------------------------------------------------------------
# E. Owner approval
# ---------------------------------------------------------------------------

def approve_submission(
    db: Session,
    app_id: uuid.UUID,
    action: str,
    current_user: User,
) -> CandidateApplication:
    if normalize_role(current_user.role.name) != ROLE_OWNER:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Owner only")

    app = _load_application(db, app_id)
    app.owner_status = action
    db.commit()
    return _load_application_with_full_relations(db, app.id)


# ---------------------------------------------------------------------------
# F. Owner dashboard — grouped by client > job
# ---------------------------------------------------------------------------

def get_owner_dashboard(db: Session, current_user: User) -> list[dict]:
    if normalize_role(current_user.role.name) != ROLE_OWNER:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Owner only")

    applications = (
        db.query(CandidateApplication)
        .options(
            joinedload(CandidateApplication.current_stage_rel),
            joinedload(CandidateApplication.submitter),
        )
        .order_by(CandidateApplication.submitted_at.desc())
        .all()
    )

    jr_ids = {a.job_requirement_id for a in applications}
    cand_ids = {a.candidate_id for a in applications}

    jr_map: dict[uuid.UUID, JobRequirement] = {}
    if jr_ids:
        for jr in (
            db.query(JobRequirement)
            .options(joinedload(JobRequirement.client))
            .filter(JobRequirement.id.in_(jr_ids))
            .all()
        ):
            jr_map[jr.id] = jr

    cand_map: dict[uuid.UUID, Candidate] = {}
    if cand_ids:
        for c in db.query(Candidate).filter(Candidate.id.in_(cand_ids)).all():
            cand_map[c.id] = c

    by_client: dict[uuid.UUID, dict] = {}
    job_to_apps: dict[uuid.UUID, list] = defaultdict(list)
    job_to_jr: dict[uuid.UUID, JobRequirement] = {}

    for app in applications:
        jr = jr_map.get(app.job_requirement_id)
        if not jr:
            continue
        job_to_apps[app.job_requirement_id].append((app, jr, cand_map.get(app.candidate_id)))
        job_to_jr[app.job_requirement_id] = jr

        if jr.client_id not in by_client:
            by_client[jr.client_id] = {"client": jr.client, "jobs": {}}
        by_client[jr.client_id]["jobs"][app.job_requirement_id] = jr

    result = []
    for client_id, client_data in by_client.items():
        client = client_data["client"]
        jobs_out = []
        for jr_id, jr in client_data["jobs"].items():
            submissions_out = [
                serialize_submission(app, jr, candidate)
                for app, _, candidate in job_to_apps.get(jr_id, [])
            ]
            jobs_out.append({
                "job_requirement_id": jr.id,
                "job_title": jr.job_title,
                "status": jr.status,
                "open_positions": jr.open_positions,
                "submissions": submissions_out,
            })
        result.append({
            "client_id": client.id,
            "company_name": client.company_name,
            "jobs": jobs_out,
        })
    return result


# ---------------------------------------------------------------------------
# G. Download — Excel in client's submission format
# ---------------------------------------------------------------------------

def generate_submission_excel(
    db: Session,
    app_id: uuid.UUID,
    current_user: User,
) -> tuple[bytes, str]:
    """Return (excel_bytes, filename)."""
    if normalize_role(current_user.role.name) != ROLE_OWNER:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Owner only")

    app = _load_application(db, app_id)
    jr = _load_job_requirement(db, app.job_requirement_id)
    candidate = _load_candidate(db, app.candidate_id)

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="openpyxl is not installed",
        )

    stored_data = _parse_stored_submission_data(app.submission_data)
    fields = _parse_submission_fields(jr.client.submission_format if jr.client else None)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Candidate Profile"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4C1D95", end_color="4C1D95", fill_type="solid")

    headers = [f.get("field", "") for f in fields]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill

    # Prefer recruiter-filled submission_data; fall back to Candidate record
    for col_idx, field_def in enumerate(fields, start=1):
        label = field_def.get("field", "")
        value = _cell_value(stored_data, label, candidate)
        if value is not None:
            ws.cell(row=2, column=col_idx, value=value)

    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    safe_name = (
        f"{candidate.first_name}_{candidate.last_name}_"
        f"{jr.client.company_name if jr.client else 'client'}"
    ).replace(" ", "_")
    return buf.read(), f"{safe_name}_profile.xlsx"


# ---------------------------------------------------------------------------
# H. Bulk download — all candidates for a job in one Excel sheet
# ---------------------------------------------------------------------------

def generate_bulk_job_excel(
    db: Session,
    job_req_id: uuid.UUID,
    current_user: User,
) -> tuple[bytes, str]:
    """Return (excel_bytes, filename) with one row per submission for the job.

    Access: owners can download any job; recruiters can only download their assigned job.
    """
    role = normalize_role(current_user.role.name)
    jr = _load_job_requirement(db, job_req_id)
    if role == ROLE_RECRUITER and jr.assigned_to != current_user.id:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Not assigned to this job")

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="openpyxl is not installed",
        )

    fields = _parse_submission_fields(jr.client.submission_format if jr.client else None)

    # Load all submissions for this job
    applications = (
        db.query(CandidateApplication)
        .filter(CandidateApplication.job_requirement_id == job_req_id)
        .order_by(CandidateApplication.submitted_at.asc())
        .all()
    )

    cand_ids = [a.candidate_id for a in applications]
    cand_map: dict[uuid.UUID, Candidate] = {}
    if cand_ids:
        for c in db.query(Candidate).filter(Candidate.id.in_(cand_ids)).all():
            cand_map[c.id] = c

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Candidates"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4C1D95", end_color="4C1D95", fill_type="solid")

    headers = [f.get("field", "") for f in fields]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill

    for row_idx, app in enumerate(applications, start=2):
        stored_data = _parse_stored_submission_data(app.submission_data)
        candidate = cand_map.get(app.candidate_id)

        for col_idx, field_def in enumerate(fields, start=1):
            label = field_def.get("field", "")
            value = _cell_value(stored_data, label, candidate)
            if value is not None:
                ws.cell(row=row_idx, column=col_idx, value=value)

    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    safe_job = (jr.job_title or "Job").replace(" ", "_")
    safe_client = (jr.client.company_name if jr.client else "Client").replace(" ", "_")
    return buf.read(), f"{safe_client}_{safe_job}_candidates.xlsx"


# ---------------------------------------------------------------------------
# I. Bulk download — all approved candidates for a client in one Excel sheet
# ---------------------------------------------------------------------------

def generate_bulk_approved_client_excel(
    db: Session,
    client_id: uuid.UUID,
    current_user: User,
) -> tuple[bytes, str]:
    """Return (excel_bytes, filename) with one row per approved submission for the client.

    Columns are taken from the client's submission_format template.
    A leading 'Job Title' column is added so the owner can see which role each row belongs to.
    """
    if normalize_role(current_user.role.name) != ROLE_OWNER:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Owner only")

    client = db.query(Client).filter(Client.id == client_id).first()
    if not client:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Client not found")

    # Collect job requirement IDs for this client as a plain Python list to
    # avoid SQLAlchemy 2.0 subquery coercion issues with .in_()
    jr_ids_list: list[uuid.UUID] = [
        row[0]
        for row in db.query(JobRequirement.id)
        .filter(JobRequirement.client_id == client_id)
        .all()
    ]

    if not jr_ids_list:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="No jobs found for this client",
        )

    applications = (
        db.query(CandidateApplication)
        .filter(
            CandidateApplication.job_requirement_id.in_(jr_ids_list),
            CandidateApplication.owner_status == "approved",
        )
        .order_by(CandidateApplication.submitted_at.asc())
        .all()
    )

    if not applications:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="No approved profiles found for this client",
        )

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="openpyxl is not installed",
        )

    fields = _parse_submission_fields(client.submission_format)

    # Load related job requirements and candidates
    jr_ids = {a.job_requirement_id for a in applications}
    cand_ids = {a.candidate_id for a in applications}

    jr_map: dict[uuid.UUID, JobRequirement] = {}
    for jr in db.query(JobRequirement).filter(JobRequirement.id.in_(jr_ids)).all():
        jr_map[jr.id] = jr

    cand_map: dict[uuid.UUID, Candidate] = {}
    for c in db.query(Candidate).filter(Candidate.id.in_(cand_ids)).all():
        cand_map[c.id] = c

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Approved Profiles"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4C1D95", end_color="4C1D95", fill_type="solid")

    # Job Title prefix column + client template fields
    all_headers = ["Job Title"] + [f.get("field", "") for f in fields]
    for col_idx, header in enumerate(all_headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill

    for row_idx, app in enumerate(applications, start=2):
        jr = jr_map.get(app.job_requirement_id)
        stored_data = _parse_stored_submission_data(app.submission_data)
        candidate = cand_map.get(app.candidate_id)

        ws.cell(row=row_idx, column=1, value=jr.job_title if jr else "")

        for col_idx, field_def in enumerate(fields, start=2):
            label = field_def.get("field", "")
            value = _cell_value(stored_data, label, candidate)
            if value is not None:
                ws.cell(row=row_idx, column=col_idx, value=value)

    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    safe_client = client.company_name.replace(" ", "_")
    return buf.read(), f"{safe_client}_approved_profiles.xlsx"


# ---------------------------------------------------------------------------
# Serializer
# ---------------------------------------------------------------------------

def serialize_submission(
    app: CandidateApplication,
    jr: JobRequirement,
    candidate: Candidate | None,
) -> dict:
    submitter_name: str | None = None
    if app.submitter:
        submitter_name = f"{app.submitter.first_name} {app.submitter.last_name}".strip()

    stage_name: str | None = None
    if app.current_stage_rel:
        stage_name = app.current_stage_rel.name

    stored_data: dict | None = None
    if app.submission_data:
        try:
            stored_data = json.loads(app.submission_data)
        except (json.JSONDecodeError, TypeError):
            stored_data = None

    return {
        "id": app.id,
        "candidate_id": app.candidate_id,
        "candidate_name": (
            f"{candidate.first_name} {candidate.last_name}".strip() if candidate else ""
        ),
        "candidate_email": candidate.email if candidate else "",
        "job_requirement_id": app.job_requirement_id,
        "job_title": jr.job_title,
        "client_id": jr.client_id,
        "client_name": jr.client.company_name if jr.client else "",
        "submitted_by_id": app.submitted_by,
        "submitted_by_name": submitter_name,
        "submitted_at": app.submitted_at,
        "current_stage": stage_name,
        "status": app.status,
        "owner_status": app.owner_status,
        "submission_data": stored_data,
    }

